"""End-to-end pipeline tests using the real parser, validator, and PPTX writer.

Only the two outbound HTTP calls are stubbed. A topic that fails drafting must
still produce a placeholder slide, and the diagnostics must account for it.
"""

from __future__ import annotations

import json

import httpx
import pytest

from app.models import Depth, Grounding
from app.services.llm import LLMService
from app.services.pipeline import Pipeline, PipelineError
from app.services.pptx_service import build_presentation
from app.services.research import ResearchService
from tests.conftest import FakeLLMTransport, llm_response, make_settings


def content_for(topic: str) -> dict:
    return {
        "slide_title": topic,
        "bullet_points": [f"{topic} point one", f"{topic} point two", f"{topic} point three"],
        "speaker_notes": f"Teaching notes for {topic}.",
    }


def build_pipeline(responses, research_transport=None, **overrides):
    settings = make_settings(**overrides)
    llm_client = httpx.AsyncClient(transport=FakeLLMTransport(responses))
    llm = LLMService(settings, client=llm_client)

    if research_transport is None:
        research = ResearchService(settings)
    else:
        research_client = httpx.AsyncClient(transport=research_transport)
        research = ResearchService(settings, client=research_client)

    return Pipeline(settings, llm=llm, research=research)


@pytest.mark.asyncio
async def test_happy_path_three_topics():
    responses = [llm_response(content_for(f"Topic {i}")) for i in range(1, 4)]
    pipeline = build_pipeline(responses)
    title, slides, diag = await pipeline.generate(
        "1. Topic 1\n2. Topic 2\n3. Topic 3", depth=Depth.HIGH_LEVEL
    )
    assert title == "Course Presentation"
    assert len(slides) == 3
    assert diag.slides_generated == 3
    assert diag.degraded_slides == 0
    assert diag.llm_calls == 3
    assert [s.slide_title for s in slides] == ["Topic 1", "Topic 2", "Topic 3"]


@pytest.mark.asyncio
async def test_empty_input_raises_pipeline_error():
    pipeline = build_pipeline([llm_response(content_for("x"))])
    with pytest.raises(PipelineError):
        await pipeline.generate("   \n  ", depth=Depth.HIGH_LEVEL)


@pytest.mark.asyncio
async def test_one_failing_topic_degrades_to_placeholder():
    responses = [
        llm_response(content_for("Topic 1")),
        httpx.Response(500, text="upstream down"),
        llm_response(content_for("Topic 3")),
    ]
    pipeline = build_pipeline(responses, max_topic_attempts=1)
    _, slides, diag = await pipeline.generate(
        "1. Topic 1\n2. Topic 2\n3. Topic 3", depth=Depth.HIGH_LEVEL
    )
    assert len(slides) == 3
    assert diag.degraded_slides == 1
    assert diag.failed_topics == ["Topic 2"]
    assert slides[1].degraded is True
    assert slides[1].grounding is Grounding.PLACEHOLDER
    assert slides[0].degraded is False
    assert slides[2].degraded is False


@pytest.mark.asyncio
async def test_research_grounding_is_recorded():
    class SearchTransport(httpx.AsyncBaseTransport):
        async def handle_async_request(self, request):
            return httpx.Response(
                200,
                json={
                    "answer": "A grounded fact.",
                    "results": [{"content": "More detail.", "url": "https://example.org/a"}],
                },
                request=request,
            )

    responses = [llm_response(content_for("Topic 1"))]
    pipeline = build_pipeline(responses, research_transport=SearchTransport(), tavily_api_key="tk")
    _, slides, diag = await pipeline.generate("1. Topic 1", depth=Depth.HIGH_LEVEL)
    assert diag.research_used is True
    assert slides[0].grounding is Grounding.RESEARCH
    assert slides[0].sources == ["https://example.org/a"]


@pytest.mark.asyncio
async def test_research_failure_degrades_but_still_generates():
    class FailingSearch(httpx.AsyncBaseTransport):
        async def handle_async_request(self, request):
            return httpx.Response(429, text="rate limited", request=request)

    responses = [llm_response(content_for("Topic 1"))]
    pipeline = build_pipeline(responses, research_transport=FailingSearch(), tavily_api_key="tk")
    _, slides, diag = await pipeline.generate("1. Topic 1", depth=Depth.HIGH_LEVEL)
    assert len(slides) == 1
    assert slides[0].grounding is Grounding.MODEL_KNOWLEDGE
    assert diag.research_used is False
    assert any("Web research is disabled" in w or "research" in w.lower() for w in diag.warnings)


@pytest.mark.asyncio
async def test_progress_callback_receives_all_stages():
    responses = [llm_response(content_for("Topic 1")), llm_response(content_for("Topic 2"))]
    pipeline = build_pipeline(responses)
    events: list[tuple[str, int, int]] = []

    async def on_progress(stage, message, completed, total):
        events.append((stage, completed, total))

    await pipeline.generate("1. Topic 1\n2. Topic 2", depth=Depth.HIGH_LEVEL, on_progress=on_progress)
    stages = [e[0] for e in events]
    assert "parsing" in stages
    assert "researching" in stages
    assert "drafting" in stages
    assert stages[-1] == "compiling"


@pytest.mark.asyncio
async def test_topic_cap_truncates_and_warns():
    responses = [llm_response(content_for(f"T{i}")) for i in range(1, 6)]
    pipeline = build_pipeline(responses, max_topics=3)
    _, slides, diag = await pipeline.generate(
        "\n".join(f"{i}. T{i}" for i in range(1, 6)), depth=Depth.HIGH_LEVEL
    )
    assert len(slides) == 3
    assert any("only the first 3" in w for w in diag.warnings)


@pytest.mark.asyncio
async def test_unconfigured_llm_raises_clear_error():
    pipeline = build_pipeline([llm_response(content_for("x"))], openrouter_api_key="")
    with pytest.raises(PipelineError) as exc:
        await pipeline.generate("1. Topic 1", depth=Depth.HIGH_LEVEL)
    assert "OPENROUTER_API_KEY" in str(exc.value)


@pytest.mark.asyncio
async def test_generated_deck_is_valid_pptx():
    responses = [llm_response(content_for(f"Topic {i}")) for i in range(1, 4)]
    pipeline = build_pipeline(responses)
    title, slides, _ = await pipeline.generate(
        "1. Topic 1\n2. Topic 2\n3. Topic 3", depth=Depth.IN_DEPTH
    )
    deck = build_presentation(slides, deck_title=title)

    from pptx import Presentation
    from io import BytesIO

    assert deck[:2] == b"PK"
    prs = Presentation(BytesIO(deck))
    # Title slide + 3 content slides + sources slide.
    assert len(prs.slides.__iter__.__self__._sldIdLst) == 5
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    joined = "\n".join(texts)
    assert "Topic 1" in joined
    assert "Sources & Method" in joined
